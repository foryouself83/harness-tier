"""Markdown to .docx for the /design-* skills.

The Markdown is the source and this only lays it out. Diagram source is POSTed to a
Kroki server, so the design leaves the host — which is why the server is config.
"""

from __future__ import annotations

import argparse
import datetime
import hashlib
import io
import re
import struct
import sys
import urllib.request
import zlib
from pathlib import Path

# Fetch alias runs eagerly at import time; collections.abc.Callable isn't subscriptable on 3.8.
from typing import Callable  # noqa: UP035
from urllib.parse import urlsplit

try:
    from _harness_paths import force_utf8_io, host_root
except ImportError:
    from scripts._harness_paths import force_utf8_io, host_root

try:
    from _design_md import ANCHOR_RE, parse, split_front
except ImportError:
    from scripts._design_md import ANCHOR_RE, parse, split_front

try:
    from design_doc_check import (
        CONVERT_ONLY,
        DOCS,
        LABEL_KEYS,
        load_settings,
        load_templates,
        source_files,
    )
except ImportError:
    from scripts.design_doc_check import (
        CONVERT_ONLY,
        DOCS,
        LABEL_KEYS,
        load_settings,
        load_templates,
        source_files,
    )

REQUIRED_MODULES = (("docx", "python-docx"), ("markdown_it", "markdown-it-py"))
# Text for design_doc_check.LABEL_KEYS, the checker's source of truth for the key set.
_LABEL_TEXT = {
    "revisions": "Revision History",
    "toc": "Contents",
    "version": "Version",
    "date": "Date",
    "summary": "Summary",
}
DEFAULT_LABELS = {k: _LABEL_TEXT[k] for k in LABEL_KEYS}
KROKI_TIMEOUT_SECONDS = 30
HEADER_FILL = "D9E2F3"
Fetch = Callable[[str, str, str], bytes]


def missing_deps() -> list[str]:
    missing = []
    for module, package in REQUIRED_MODULES:
        try:
            __import__(module)
        except ImportError:
            missing.append(package)
    return missing


def kroki_fetch(url: str, engine: str, source: str) -> bytes:
    # png is 500/unsupported on some engines (measured against kroki.io); svg is universal.
    req = urllib.request.Request(
        f"{url.rstrip('/')}/{engine}/svg",
        data=source.encode("utf-8"),
        headers={"Content-Type": "text/plain", "User-Agent": "harness-tier-design-docs"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=KROKI_TIMEOUT_SECONDS) as resp:
        return resp.read()


def bookmark_name(anchor: str) -> str:
    # Word drops a bookmark whose name has a hyphen or runs past 40 characters.
    name = "a_" + anchor.lower().replace("-", "_")
    if len(name) <= 40:
        return name
    # Two anchors sharing a 40-char prefix would otherwise collide once truncated.
    digest = hashlib.sha1(anchor.encode("utf-8")).hexdigest()[:8]
    return f"{name[:31]}_{digest}"


_SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
_SVG_EXT_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"


def _svg_aspect(svg: bytes) -> float:
    """Height/width from the SVG's viewBox or width/height attributes; 16:9 when neither parses."""
    text = svg.decode("utf-8", errors="ignore")
    m = re.search(r'viewBox\s*=\s*"([^"]+)"', text)
    if m:
        nums = re.findall(r"[-+]?[0-9.]+(?:[eE][-+]?[0-9]+)?", m.group(1))
        if len(nums) == 4 and float(nums[2]) > 0:
            return float(nums[3]) / float(nums[2])
    w = re.search(r'\bwidth\s*=\s*"([0-9.]+)', text)
    h = re.search(r'\bheight\s*=\s*"([0-9.]+)', text)
    if w and h and float(w.group(1)) > 0:
        return float(h.group(1)) / float(w.group(1))
    return 9 / 16


def _blank_png() -> bytes:
    """A minimal 1x1 raster fallback for pre-2016 Word, which ignores the SVG extension."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        crc = zlib.crc32(kind + data)
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", crc)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00\xff\xff\xff")
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


_FALLBACK_PNG = _blank_png()


def _embed_svg(doc_part, shape, svg: bytes) -> None:
    """Word 2016+ renders the vector via an asvg:svgBlip on the fallback picture's blip."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.part import Part
    from docx.oxml.parser import parse_xml

    partname = doc_part.package.next_partname("/word/media/image%d.svg")
    svg_part = Part(partname, "image/svg+xml", svg, doc_part.package)
    rid = doc_part.relate_to(svg_part, RT.IMAGE)
    ext = parse_xml(
        '<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:ext uri="{_SVG_EXT_URI}" xmlns:asvg="{_SVG_EXT_NS}">'
        "<asvg:svgBlip"
        ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        f' r:embed="{rid}"/></a:ext></a:extLst>'
    )
    shape._inline.graphic.graphicData.pic.blipFill.blip.append(ext)


class _Writer:
    def __init__(self, document, fetch: Fetch, url: str, anchors: set[str]):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt

        self.d = document
        self.fetch = fetch
        self.url = url
        self.anchors = anchors
        self.warnings: list[str] = []
        self._next_id = 0
        self._el, self._qn = OxmlElement, qn
        self._inches, self._pt = Inches, Pt
        self._url_scheme = urlsplit(url).scheme

    def bookmark(self, paragraph, anchor: str) -> None:
        start = self._el("w:bookmarkStart")
        start.set(self._qn("w:id"), str(self._next_id))
        start.set(self._qn("w:name"), bookmark_name(anchor))
        end = self._el("w:bookmarkEnd")
        end.set(self._qn("w:id"), str(self._next_id))
        paragraph._p.append(start)
        paragraph._p.append(end)
        self._next_id += 1

    def hyperlink(self, paragraph, text: str, anchor: str) -> None:
        link = self._el("w:hyperlink")
        link.set(self._qn("w:anchor"), bookmark_name(anchor))
        run = self._el("w:r")
        props = self._el("w:rPr")
        color = self._el("w:color")
        color.set(self._qn("w:val"), "0563C1")
        underline = self._el("w:u")
        underline.set(self._qn("w:val"), "single")
        props.append(color)
        props.append(underline)
        run.append(props)
        t = self._el("w:t")
        t.text = text
        t.set(self._qn("xml:space"), "preserve")
        run.append(t)
        link.append(run)
        paragraph._p.append(link)

    def _target(self, href: str) -> str | None:
        frag = href.split("#", 1)[1].lower() if "#" in href else ""
        return frag if frag and frag in self.anchors else None

    def inline(self, paragraph, children) -> None:
        bold = italic = False
        target: str | None = None
        for tok in children or []:
            kind = tok.type
            if kind == "strong_open":
                bold = True
            elif kind == "strong_close":
                bold = False
            elif kind == "em_open":
                italic = True
            elif kind == "em_close":
                italic = False
            elif kind == "link_open":
                target = self._target(tok.attrGet("href") or "")
            elif kind == "link_close":
                target = None
            elif kind == "html_inline":
                for m in ANCHOR_RE.finditer(tok.content):
                    self.bookmark(paragraph, m.group(1))
            elif kind in ("text", "code_inline"):
                if target:
                    self.hyperlink(paragraph, tok.content, target)
                    continue
                run = paragraph.add_run(tok.content)
                run.bold, run.italic = bold, italic
                if kind == "code_inline":
                    run.font.name = "Consolas"
            elif kind == "softbreak":
                paragraph.add_run(" ")
            elif kind == "hardbreak":
                paragraph.add_run().add_break()

    def _paragraph(self, style: str | None):
        try:
            return self.d.add_paragraph(style=style)
        except KeyError:  # a base_docx without list styles
            return self.d.add_paragraph()

    def _shade(self, cell) -> None:
        shd = self._el("w:shd")
        shd.set(self._qn("w:val"), "clear")
        shd.set(self._qn("w:color"), "auto")
        shd.set(self._qn("w:fill"), HEADER_FILL)
        cell._tc.get_or_add_tcPr().append(shd)

    def _table(self, tokens, i: int) -> int:
        rows: list[list] = []
        while tokens[i].type != "table_close":
            tok = tokens[i]
            if tok.type == "tr_open":
                rows.append([])
            elif tok.type == "inline":
                rows[-1].append(tok.children)
            i += 1
        ncols = max((len(r) for r in rows), default=0)
        if not ncols:
            return i + 1
        table = self.d.add_table(rows=len(rows), cols=ncols)
        try:
            table.style = "Table Grid"
        except (KeyError, ValueError):
            pass
        for r, row in enumerate(rows):
            for c, children in enumerate(row[:ncols]):
                cell = table.cell(r, c)
                self.inline(cell.paragraphs[0], children)
                if r == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                    self._shade(cell)
        return i + 1

    def code(self, text: str) -> None:
        run = self.d.add_paragraph().add_run(text.rstrip("\n"))
        run.font.name = "Consolas"
        run.font.size = self._pt(9)

    def fence(self, lang: str, body: str, where: str) -> None:
        if lang in ("mermaid", "d2"):
            if self._url_scheme not in ("http", "https"):
                self.warnings.append(
                    f"{lang} diagram in {where} not rendered (renderer URL scheme "
                    f"'{self._url_scheme or self.url}' is not http/https); source kept as code"
                )
                self.code(body)
                return
            try:
                svg = self.fetch(self.url, lang, body)
                shape = self.d.add_picture(
                    io.BytesIO(_FALLBACK_PNG),
                    width=self._inches(6.0),
                    height=self._inches(6.0 * _svg_aspect(svg)),
                )
                _embed_svg(self.d.part, shape, svg)
                return
            except Exception as exc:  # any renderer or embedding failure keeps the source instead
                self.warnings.append(
                    f"{lang} diagram in {where} not rendered ({exc}); source kept as code"
                )
        self.code(body)

    def markdown(self, body: str, where: str) -> None:
        from markdown_it import MarkdownIt

        tokens = MarkdownIt("commonmark", {"html": True}).enable("table").parse(body)
        lists: list[str] = []
        i = 0
        while i < len(tokens):
            tok = tokens[i]
            kind = tok.type
            if kind == "heading_open":
                p = self.d.add_heading("", level=min(int(tok.tag[1]), 9))
                self.inline(p, tokens[i + 1].children)
                i += 3
                continue
            if kind == "paragraph_open":
                style = None
                if lists:
                    base = "List Bullet" if lists[-1] == "bullet" else "List Number"
                    style = base if len(lists) == 1 else f"{base} {min(len(lists), 3)}"
                self.inline(self._paragraph(style), tokens[i + 1].children)
                i += 3
                continue
            if kind in ("bullet_list_open", "ordered_list_open"):
                lists.append("bullet" if kind.startswith("bullet") else "ordered")
            elif kind in ("bullet_list_close", "ordered_list_close"):
                lists.pop()
            elif kind == "table_open":
                i = self._table(tokens, i)
                continue
            elif kind in ("fence", "code_block"):
                self.fence((tok.info or "").strip().lower(), tok.content, where)
            elif kind == "html_block":
                for m in ANCHOR_RE.finditer(tok.content):
                    self.bookmark(self.d.add_paragraph(), m.group(1))
            i += 1


def _front_pages(d, title: str, standard: str, labels: dict[str, str], revisions: list) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    version = str(revisions[-1].get("version", "")) if revisions else ""
    lines = [
        (title, 28, True),
        (standard, 14, False),
        (f"{labels['version']} {version}  ·  {datetime.date.today().isoformat()}", 12, False),
    ]
    for text, size, bold in lines:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size, run.bold = Pt(size), bold
    d.add_page_break()
    d.add_paragraph(labels["revisions"]).runs[0].bold = True
    table = d.add_table(rows=1, cols=3)
    for cell, key in zip(table.rows[0].cells, ("version", "date", "summary")):
        cell.text = labels[key]
    for rev in revisions:
        cells = table.add_row().cells
        for cell, key in zip(cells, ("version", "date", "summary")):
            cell.text = str(rev.get(key, "")) if isinstance(rev, dict) else ""
    d.add_page_break()
    d.add_paragraph(labels["toc"]).runs[0].bold = True
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    d.add_paragraph()._p.append(field)
    d.add_page_break()
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    d.settings.element.append(update)


def render(
    root: Path, settings: dict, name: str, fetch: Fetch = kroki_fetch
) -> tuple[Path, list[str]]:
    from docx import Document

    template = load_templates(root / settings["templates"])[name]
    tfront = template.front or {}
    if name in CONVERT_ONLY:
        files = source_files(root, tfront.get("sources") or [])
        revisions = tfront.get("revisions") or []
    else:
        files = [root / settings["docs"] / f"{name}.md"]
        if not files[0].is_file():
            rel = files[0].relative_to(root).as_posix()
            raise FileNotFoundError(f"{rel} not found — write it first with /design-{name}")
        front, _, _, _ = split_front(files[0].read_text(encoding="utf-8"))
        revisions = (front or {}).get("revisions") or []
    texts = [(p, split_front(p.read_text(encoding="utf-8"))[2]) for p in files]
    anchors: set[str] = set()
    for _, body in texts:
        anchors |= {a for a, _ in parse(body).anchors}
    base = settings.get("base_docx")
    d = Document(str(root / base)) if base else Document()
    labels = dict(DEFAULT_LABELS)
    labels.update({k: str(v) for k, v in (tfront.get("labels") or {}).items()})
    _front_pages(
        d, str(tfront.get("title", name)), str(tfront.get("standard", "")), labels, revisions
    )
    writer = _Writer(d, fetch, str(settings["renderer"]), anchors)
    for k, (path, body) in enumerate(texts):
        if k:
            d.add_page_break()
        writer.markdown(body, path.relative_to(root).as_posix())
    out = root / settings["output"] / f"{name}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
    return out, writer.warnings


def main(argv: list[str] | None = None) -> int:
    force_utf8_io()
    parser = argparse.ArgumentParser(description="Render a design document to .docx.")
    parser.add_argument("--root", default=None, help="host root (default: project dir)")
    parser.add_argument("--check-deps", action="store_true", help="report missing Python packages")
    parser.add_argument("doc", nargs="?", choices=DOCS)
    args = parser.parse_args(argv)
    missing = missing_deps()
    if args.check_deps or missing:
        if missing:
            print(f"missing: {' '.join(missing)}")
            return 3
        print("ok")
        return 0
    if not args.doc:
        parser.error("doc is required")
    root = Path(args.root) if args.root else host_root()
    try:
        out, warnings = render(root, load_settings(root), args.doc)
    except FileNotFoundError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    for w in warnings:
        print(f"warning: {w}")
    print(out.relative_to(root).as_posix())
    return 0


if __name__ == "__main__":
    sys.exit(main())
