"""The docx is read back, not eyeballed: headings, tables, bookmarks and pictures must
follow the Markdown and the consumer's template."""

from docx import Document
from docx.oxml.ns import qn

from scripts.design_doc_check import load_settings
from scripts.design_doc_render import _SVG_EXT_NS, bookmark_name, kroki_fetch, main, render
from tests.design_docs._fixture import TDIR, TEMPLATES, make_host
from tests.design_docs._helpers import write

_SVG = (
    b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100">'
    b'<rect width="200" height="100"/></svg>'
)
_SVGBLIP_TAG = f"{{{_SVG_EXT_NS}}}svgBlip"


def ok_fetch(url, engine, source):
    ok_fetch.calls.append((url, engine))
    return _SVG


def bad_fetch(url, engine, source):
    raise OSError("offline")


def headings(doc):
    return [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]


def test_deliverable_renders_structure(tmp_path):
    make_host(tmp_path)
    out, warnings = render(tmp_path, load_settings(tmp_path), "table", fetch=ok_fetch)
    assert out == tmp_path / "docs/deliverables/results/table.docx" and warnings == []
    doc = Document(str(out))
    assert "2. 테이블 상세" in headings(doc)
    assert any("TBL-001 payment" in h for h in headings(doc))
    headers = [[c.text for c in t.rows[0].cells] for t in doc.tables]
    assert ["컬럼명", "타입", "FK"] in headers
    marks = [b.get(qn("w:name")) for b in doc.element.body.iter(qn("w:bookmarkStart"))]
    assert bookmark_name("tbl-001") in marks
    xml = doc.settings.element.xml
    assert "updateFields" in xml


ok_fetch.calls = []


def test_revision_table_and_cover(tmp_path):
    make_host(tmp_path)
    out, _ = render(tmp_path, load_settings(tmp_path), "erd", fetch=ok_fetch)
    doc = Document(str(out))
    assert doc.paragraphs[0].text == "ERD"
    rows = [[c.text for c in r.cells] for t in doc.tables for r in t.rows]
    assert ["1.0", "2026-09-21", "first"] in rows


def test_internal_link_becomes_hyperlink(tmp_path):
    make_host(tmp_path)
    out, _ = render(tmp_path, load_settings(tmp_path), "architecture", fetch=ok_fetch)
    anchors = [
        h.get(qn("w:anchor")) for h in Document(str(out)).element.body.iter(qn("w:hyperlink"))
    ]
    assert bookmark_name("cmp-001") in anchors


def test_diagram_picture_and_fallback(tmp_path):
    make_host(tmp_path)
    ok_fetch.calls.clear()
    out, warnings = render(tmp_path, load_settings(tmp_path), "sds", fetch=ok_fetch)
    assert ok_fetch.calls == [("https://kroki.io", "mermaid")] and warnings == []
    doc = Document(str(out))
    assert len(doc.inline_shapes) == 1
    shape = doc.inline_shapes[0]
    assert abs(shape.height / shape.width - 0.5) <= 0.005
    assert next(doc.element.body.iter(_SVGBLIP_TAG), None) is not None
    svg_parts = [p for p in doc.part.package.iter_parts() if p.content_type == "image/svg+xml"]
    assert len(svg_parts) == 1

    out, warnings = render(tmp_path, load_settings(tmp_path), "sds", fetch=bad_fetch)
    assert len(warnings) == 1 and "mermaid" in warnings[0]
    assert any("graph TD" in p.text for p in Document(str(out)).paragraphs)


def test_srs_merges_sources_in_template_order(tmp_path):
    make_host(tmp_path)
    out, _ = render(tmp_path, load_settings(tmp_path), "srs", fetch=ok_fetch)
    hs = headings(Document(str(out)))
    assert hs.index("SRS") < hs.index("Pay")


def test_template_labels_and_output_setting(tmp_path):
    make_host(tmp_path)
    write(
        tmp_path,
        f"{TDIR}/erd.template.md",
        TEMPLATES["erd"].replace("inventory:", "labels: {revisions: 개정 이력}\ninventory:"),
    )
    write(tmp_path, ".claude/harness-tier/config/flow-config.yaml", "design_docs:\n  output: out\n")
    out, _ = render(tmp_path, load_settings(tmp_path), "erd", fetch=ok_fetch)
    assert out == tmp_path / "out/erd.docx"
    assert any(p.text == "개정 이력" for p in Document(str(out)).paragraphs)


def test_cli_reports_missing_deps(tmp_path, monkeypatch, capsys):
    import scripts.design_doc_render as r

    monkeypatch.setattr(r, "missing_deps", lambda: ["python-docx"])
    assert main(["--root", str(tmp_path), "--check-deps"]) == 3
    assert "python-docx" in capsys.readouterr().out


def test_kroki_fetch_posts_svg_with_user_agent(monkeypatch):
    captured = {}

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def read(self):
            return b"<svg/>"

    def fake_urlopen(req, timeout=None):
        captured["url"] = req.full_url
        captured["timeout"] = timeout
        captured["headers"] = {k.lower(): v for k, v in req.header_items()}
        return FakeResponse()

    monkeypatch.setattr("urllib.request.urlopen", fake_urlopen)
    out = kroki_fetch("https://kroki.io", "mermaid", "graph TD\n  A --> B")
    assert out == b"<svg/>"
    assert captured["url"] == "https://kroki.io/mermaid/svg"
    assert captured["headers"]["user-agent"] == "harness-tier-design-docs"


def test_srs_render_handles_a_template_with_no_revisions(tmp_path):
    make_host(tmp_path)
    out, _ = render(tmp_path, load_settings(tmp_path), "srs", fetch=ok_fetch)
    doc = Document(str(out))
    assert doc.paragraphs[2].text.startswith("Version")
    assert "1.0" not in doc.paragraphs[2].text


def test_render_missing_document_reports_a_clear_error(tmp_path, capsys):
    make_host(tmp_path)
    (tmp_path / "docs/deliverables/erd.md").unlink()
    assert main(["--root", str(tmp_path), "erd"]) == 1
    err = capsys.readouterr().err
    assert "docs/deliverables/erd.md" in err and "/design-erd" in err


def test_diagram_falls_back_when_the_renderer_url_scheme_is_not_http(tmp_path):
    make_host(tmp_path)
    write(
        tmp_path,
        ".claude/harness-tier/config/flow-config.yaml",
        "design_docs:\n  renderer: ftp://kroki.local\n",
    )
    ok_fetch.calls.clear()
    out, warnings = render(tmp_path, load_settings(tmp_path), "sds", fetch=ok_fetch)
    assert ok_fetch.calls == []
    assert len(warnings) == 1 and "ftp" in warnings[0]
    doc = Document(str(out))
    assert any("graph TD" in p.text for p in doc.paragraphs)


def test_bookmark_name_collision_safe_past_forty_chars():
    prefix = "c-" * 20  # 40 identical characters once anchors are lower-cased
    anchor_a = prefix + "a" * 20
    anchor_b = prefix + "b" * 20
    name_a, name_b = bookmark_name(anchor_a), bookmark_name(anchor_b)
    assert name_a != name_b
    assert len(name_a) <= 40 and len(name_b) <= 40
